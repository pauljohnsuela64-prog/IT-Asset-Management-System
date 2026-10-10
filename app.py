from datetime import date, datetime
import os
from pathlib import Path
from dotenv import load_dotenv
from flask import Flask, render_template, request, redirect, url_for, abort
from flask import session

from werkzeug.security import (
    check_password_hash,
    generate_password_hash,
)
from functools import wraps


import mysql.connector

from io import BytesIO

from flask import send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from export_repository import (
    get_assets_for_export,
    get_employees_for_export,
    get_assignments_for_export,
    get_maintenance_for_export,
)
from report_repository import (
    get_asset_status_report,
    get_asset_type_report,
)

from dashboard_repository import get_dashboard_stats


from user_repository import (
    get_user_by_username,
    get_all_users,
    create_user,
    get_user_by_id,
    update_user_status,
)


from maintenance_repository import (
    get_all_maintenance_records,
    get_maintainable_assets,
    create_maintenance_record,
    get_maintenance_by_id,
    complete_maintenance,
)
from employee_repository import (
    get_all_employees,
    create_employee,
    search_employees,
    get_employee_by_id,
    update_employee,
    deactivate_employee,
)

from asset_repository import (
    get_all_assets,
    create_asset,
    search_assets,
    get_asset_by_id,
    update_asset,
    retire_asset,
    asset_has_history,
    delete_asset,
)

from assignment_repository import (
    get_available_assets,
    get_active_employees,
    create_assignment,
    get_all_assignments,
    get_assignment_by_id,
    return_assignment,
)


app = Flask(__name__)
load_dotenv(Path(__file__).resolve().with_name(".env"))
app.secret_key = os.getenv("SECRET_KEY")
if not app.secret_key or app.secret_key in (
    "replace_with_a_long_random_secret_key",
    "change-this-to-a-random-secret-key",
):
    raise RuntimeError("Set SECRET_KEY to a long random value in .env before starting the app.")

def login_required(view_function):
    @wraps(view_function)
    def wrapped_view(*args, **kwargs):
        if "user_id" not in session:
            return redirect(url_for("login"))

        return view_function(*args, **kwargs)

    return wrapped_view

def admin_required(view_function):
    @wraps(view_function)
    def wrapped_view(*args, **kwargs):
        if "user_id" not in session:
            return redirect(url_for("login"))

        if session.get("role") != "Admin":
            return "Admin access required.", 403

        return view_function(*args, **kwargs)

    return wrapped_view


@app.route("/")
@login_required
def home():
    search_term = request.args.get("search", "").strip()

    if search_term:
        assets = search_assets(search_term)
    else:
        assets = get_all_assets()

    return render_template(
        "assets.html",
        assets=assets,
        search_term=search_term,
    )


@app.route("/assets/add", methods=["GET", "POST"])
@login_required
def add_asset():
    error = None

    if request.method == "POST":
        asset_tag = request.form.get("asset_tag", "").strip()
        device_name = request.form.get("device_name", "").strip()
        asset_type = request.form.get("asset_type", "").strip()
        brand = request.form.get("brand", "").strip()

        model = request.form.get("model", "").strip() or None
        serial_number = request.form.get("serial_number", "").strip() or None
        purchase_date = request.form.get("purchase_date", "").strip() or None
        notes = request.form.get("notes", "").strip() or None

        if not asset_tag:
            error = "Asset Tag is required."

        elif not device_name:
            error = "Device Name is required."

        elif not asset_type:
            error = "Asset Type is required."

        elif not brand:
            error = "Brand is required."

        else:
            try:
                create_asset(
                    asset_tag,
                    device_name,
                    asset_type,
                    brand,
                    model,
                    serial_number,
                    purchase_date,
                    notes,
                )

                return redirect(url_for("home"))

            except mysql.connector.IntegrityError:
                error = "Asset Tag or Serial Number already exists."

            except mysql.connector.Error:
                error = "Unable to save the asset. Please try again."

    return render_template(
        "add_asset.html",
        error=error,
    )

@app.route("/assets/<int:asset_id>/edit", methods=["GET", "POST"])
@admin_required
def edit_asset(asset_id):
    asset = get_asset_by_id(asset_id)

    if asset is None:
        abort(404)

    if asset["status"] == "Retired":
        return "Retired assets cannot be edited.", 403

    error = None

    if request.method == "POST":
        asset_tag = request.form.get("asset_tag", "").strip()
        device_name = request.form.get("device_name", "").strip()
        asset_type = request.form.get("asset_type", "").strip()
        brand = request.form.get("brand", "").strip()

        model = request.form.get("model", "").strip() or None
        serial_number = request.form.get("serial_number", "").strip() or None
        status = request.form.get("status", "").strip()

        original_status = asset["status"]

        purchase_date = request.form.get("purchase_date", "").strip() or None
        notes = request.form.get("notes", "").strip() or None

        if not asset_tag:
            error = "Asset Tag is required."

        elif not device_name:
            error = "Device Name is required."

        elif not asset_type:
            error = "Asset Type is required."

        elif not brand:
            error = "Brand is required."

        elif original_status == "Assigned" and status != "Assigned":
            error = (
                "Assigned assets must be returned "
                "before their status can change."
            )

        elif original_status != "Assigned" and status not in (
            "Available",
            "Under Maintenance",
        ):
            error = "Invalid asset status."

        else:
            try:
                update_asset(
                    asset_id,
                    asset_tag,
                    device_name,
                    asset_type,
                    brand,
                    model,
                    serial_number,
                    status,
                    purchase_date,
                    notes,
                )

                return redirect(url_for("home"))

            except mysql.connector.IntegrityError:
                error = "Asset Tag or Serial Number already exists."

            except mysql.connector.Error:
                error = "Unable to update the asset. Please try again."

        asset = {
            "asset_id": asset_id,
            "asset_tag": asset_tag,
            "device_name": device_name,
            "asset_type": asset_type,
            "brand": brand,
            "model": model,
            "serial_number": serial_number,
            "status": status,
            "purchase_date": purchase_date,
            "notes": notes,
        }

    return render_template(
        "edit_asset.html",
        asset=asset,
        error=error,
    )


@app.route("/assets/<int:asset_id>/retire", methods=["GET", "POST"])
@admin_required
def retire_asset_route(asset_id):
    asset = get_asset_by_id(asset_id)

    if asset is None:
        abort(404)

    if asset["status"] == "Assigned":
        return (
            "Assigned assets must be returned before they can be retired.",
            400,
        )

    if asset["status"] == "Retired":
        return "This asset is already retired.", 400

    error = None

    if request.method == "POST":
        try:
            retire_asset(asset_id)

            return redirect(url_for("home"))

        except ValueError as error_message:
            error = str(error_message)

        except mysql.connector.Error:
            error = "Unable to retire the asset. Please try again."

    return render_template(
        "retire_asset.html",
        asset=asset,
        error=error,
    )

@app.route("/employees")
@login_required
def employees():
    search_term = request.args.get("search", "").strip()

    if search_term:
        employees = search_employees(search_term)
    else:
        employees = get_all_employees()

    return render_template(
        "employees.html",
        employees=employees,
        search_term=search_term,
    )

@app.route("/employees/add", methods=["GET", "POST"])
@login_required
def add_employee():
    error = None

    if request.method == "POST":
        employee_code = request.form.get("employee_code", "").strip()
        full_name = request.form.get("full_name", "").strip()
        department = request.form.get("department", "").strip()

        position = request.form.get("position", "").strip() or None
        email = request.form.get("email", "").strip() or None

        if not employee_code:
            error = "Employee Code is required."

        elif not full_name:
            error = "Full Name is required."

        elif not department:
            error = "Department is required."

        else:
            try:
                create_employee(
                    employee_code,
                    full_name,
                    department,
                    position,
                    email,
                )

                return redirect(url_for("employees"))

            except mysql.connector.IntegrityError:
                error = "Employee Code or Email already exists."

            except mysql.connector.Error:
                error = "Unable to save the employee. Please try again."

    return render_template(
        "add_employee.html",
        error=error,
    )

@app.route("/employees/<int:employee_id>/edit", methods=["GET", "POST"])
@login_required
def edit_employee(employee_id):
    employee = get_employee_by_id(employee_id)

    if employee is None:
        abort(404)

    error = None

    if request.method == "POST":
        employee_code = request.form.get("employee_code", "").strip()
        full_name = request.form.get("full_name", "").strip()
        department = request.form.get("department", "").strip()

        position = request.form.get("position", "").strip() or None
        email = request.form.get("email", "").strip() or None
        status = request.form.get("status", "").strip()

        if not employee_code:
            error = "Employee Code is required."

        elif not full_name:
            error = "Full Name is required."

        elif not department:
            error = "Department is required."

        elif status not in ("Active", "Inactive"):
            error = "Invalid employee status."

        else:
            try:
                update_employee(
                    employee_id,
                    employee_code,
                    full_name,
                    department,
                    position,
                    email,
                    status,
                )

                return redirect(url_for("employees"))

            except mysql.connector.IntegrityError:
                error = "Employee Code or Email already exists."

            except mysql.connector.Error:
                error = "Unable to update the employee. Please try again."

        employee = {
            "employee_id": employee_id,
            "employee_code": employee_code,
            "full_name": full_name,
            "department": department,
            "position": position,
            "email": email,
            "status": status,
        }

    return render_template(
        "edit_employee.html",
        employee=employee,
        error=error,
    )


@app.route("/employees/<int:employee_id>/deactivate", methods=["GET", "POST"])
@admin_required
def deactivate_employee_route(employee_id):
    employee = get_employee_by_id(employee_id)

    if employee is None:
        abort(404)

    if request.method == "POST":
        try:
            deactivate_employee(employee_id)
            return redirect(url_for("employees"))

        except mysql.connector.Error:
            return render_template(
                "deactivate_employee.html",
                employee=employee,
                error="Unable to deactivate the employee. Please try again.",
            )

    return render_template(
        "deactivate_employee.html",
        employee=employee,
        error=None,
    )

@app.route("/assignments/add", methods=["GET", "POST"])
@login_required
def add_assignment():
    available_assets = get_available_assets()
    active_employees = get_active_employees()

    error = None

    if request.method == "POST":
        asset_id = request.form.get("asset_id", "").strip()
        employee_id = request.form.get("employee_id", "").strip()
        assigned_date = request.form.get("assigned_date", "").strip()
        notes = request.form.get("notes", "").strip() or None

        if not asset_id:
            error = "Please select an asset."

        elif not employee_id:
            error = "Please select an employee."

        elif not assigned_date:
            error = "Assigned Date is required."

        else:
            try:
                assigned_date_value = date.fromisoformat(assigned_date)

                if assigned_date_value > date.today():
                    error = "Assigned Date cannot be in the future."

                else:
                    create_assignment(
                        int(asset_id),
                        int(employee_id),
                        assigned_date_value,
                        notes,
                    )

                    return redirect(url_for("home"))

            except ValueError as error_message:
                error = str(error_message)

            except mysql.connector.Error:
                error = "Unable to assign the asset. Please try again."

    return render_template(
        "add_assignment.html",
        assets=available_assets,
        employees=active_employees,
        error=error,
    )


@app.route("/assignments")
@login_required
def assignments():
    assignments = get_all_assignments()

    return render_template(
        "assignments.html",
        assignments=assignments,
    )

@app.route(
    "/assignments/<int:assignment_id>/return",
    methods=["GET", "POST"],
)
@login_required
def return_asset_route(assignment_id):
    assignment = get_assignment_by_id(assignment_id)

    if assignment is None:
        abort(404)

    if assignment["status"] != "Assigned":
        return "This asset has already been returned.", 400

    error = None

    if request.method == "POST":
        returned_date = request.form.get("returned_date", "").strip()

        if not returned_date:
            error = "Returned Date is required."

        else:
            try:
                returned_date_value = date.fromisoformat(returned_date)

                assigned_date_value = assignment["assigned_date"]

                if isinstance(assigned_date_value, str):
                    assigned_date_value = date.fromisoformat(
                        assigned_date_value
                    )

                if returned_date_value < assigned_date_value:
                    error = (
                        "Returned Date cannot be earlier "
                        "than the Assigned Date."
                    )

                elif returned_date_value > date.today():
                    error = "Returned Date cannot be in the future."

                else:
                    try:
                        return_assignment(
                            assignment_id,
                            returned_date_value,
                        )

                        return redirect(url_for("assignments"))

                    except ValueError as error_message:
                        error = str(error_message)

                    except mysql.connector.Error:
                        error = (
                            "Unable to return the asset. "
                            "Please try again."
                        )

            except ValueError:
                error = "Invalid Returned Date."

    return render_template(
        "return_asset.html",
        assignment=assignment,
        error=error,
    )

@app.route("/maintenance")
@login_required
def maintenance():
    records = get_all_maintenance_records()

    return render_template(
        "maintenance.html",
        records=records,
    )

@app.route("/maintenance/add", methods=["GET", "POST"])
@login_required
def add_maintenance():
    assets = get_maintainable_assets()
    error = None

    if request.method == "POST":
        asset_id = request.form.get("asset_id", "").strip()
        maintenance_date = request.form.get("maintenance_date", "").strip()
        maintenance_type = request.form.get("maintenance_type", "").strip()
        issue_description = request.form.get("issue_description", "").strip()

        technician_vendor = (
            request.form.get("technician_vendor", "").strip() or None
        )

        cost = request.form.get("cost", "").strip() or "0"
        notes = request.form.get("notes", "").strip() or None

        if not asset_id:
            error = "Please select an asset."

        elif not maintenance_date:
            error = "Maintenance Date is required."

        elif not maintenance_type:
            error = "Maintenance Type is required."

        elif not issue_description:
            error = "Issue Description is required."

        else:
            try:
                maintenance_date_value = date.fromisoformat(
                    maintenance_date
                )

                if maintenance_date_value > date.today():
                    error = "Maintenance Date cannot be in the future."

                else:
                    cost_value = float(cost)

                    if cost_value < 0:
                        error = "Cost cannot be negative."

                    else:
                        create_maintenance_record(
                            int(asset_id),
                            maintenance_date_value,
                            maintenance_type,
                            issue_description,
                            technician_vendor,
                            cost_value,
                            notes,
                        )

                        return redirect(url_for("maintenance"))

            except ValueError as error_message:
                error = str(error_message)

            except mysql.connector.Error:
                error = (
                    "Unable to create the maintenance record. "
                    "Please try again."
                )

    return render_template(
        "add_maintenance.html",
        assets=assets,
        error=error,
    )

@app.route(
    "/maintenance/<int:maintenance_id>/complete",
    methods=["GET", "POST"],
)
@login_required
def complete_maintenance_route(maintenance_id):
    record = get_maintenance_by_id(maintenance_id)

    if record is None:
        abort(404)

    if record["status"] == "Completed":
        return "This maintenance record is already completed.", 400

    error = None

    if request.method == "POST":
        completed_date = request.form.get("completed_date", "").strip()
        action_taken = request.form.get("action_taken", "").strip()

        technician_vendor = (
            request.form.get("technician_vendor", "").strip() or None
        )

        cost = request.form.get("cost", "").strip() or "0"
        notes = request.form.get("notes", "").strip() or None

        if not completed_date:
            error = "Completed Date is required."

        elif not action_taken:
            error = "Action Taken is required."

        else:
            try:
                completed_date_value = date.fromisoformat(completed_date)

                maintenance_date_value = record["maintenance_date"]

                if isinstance(maintenance_date_value, str):
                    maintenance_date_value = date.fromisoformat(
                        maintenance_date_value
                    )

                if completed_date_value < maintenance_date_value:
                    error = (
                        "Completed Date cannot be earlier "
                        "than the Maintenance Date."
                    )

                elif completed_date_value > date.today():
                    error = "Completed Date cannot be in the future."

                else:
                    cost_value = float(cost)

                    if cost_value < 0:
                        error = "Cost cannot be negative."

                    else:
                        complete_maintenance(
                            maintenance_id,
                            completed_date_value,
                            action_taken,
                            technician_vendor,
                            cost_value,
                            notes,
                        )

                        return redirect(url_for("maintenance"))

            except ValueError as error_message:
                error = str(error_message)

            except mysql.connector.Error:
                error = (
                    "Unable to complete maintenance. "
                    "Please try again."
                )

    return render_template(
        "complete_maintenance.html",
        record=record,
        error=error,
    )


@app.route("/login", methods=["GET", "POST"])
def login():
    error = None

    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")

        if not username:
            error = "Username is required."

        elif not password:
            error = "Password is required."

        else:
            user = get_user_by_username(username)

            if user is None:
                error = "Invalid username or password."

            elif user["status"] != "Active":
                error = "This account is inactive."

            elif not check_password_hash(
                user["password_hash"],
                password,
            ):
                error = "Invalid username or password."

            else:
                session["user_id"] = user["user_id"]
                session["username"] = user["username"]
                session["role"] = user["role"]

                return redirect(url_for("home"))

    return render_template(
        "login.html",
        error=error,
    )


@app.route("/logout")
def logout():
    session.clear()

    return redirect(url_for("login"))


@app.route("/users")
@admin_required
def users():
    users = get_all_users()

    return render_template(
        "users.html",
        users=users,
    )


@app.route("/users/add", methods=["GET", "POST"])
@admin_required
def add_user():
    error = None

    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        confirm_password = request.form.get("confirm_password", "")
        role = request.form.get("role", "").strip()

        if not username:
            error = "Username is required."

        elif not password:
            error = "Password is required."

        elif password != confirm_password:
            error = "Passwords do not match."

        elif role not in ("Admin", "Staff"):
            error = "Invalid user role."

        else:
            try:
                password_hash = generate_password_hash(password)

                create_user(
                    username,
                    password_hash,
                    role,
                )

                return redirect(url_for("users"))

            except mysql.connector.IntegrityError:
                error = "Username already exists."

            except mysql.connector.Error:
                error = "Unable to create the user. Please try again."

    return render_template(
        "add_user.html",
        error=error,
    )


@app.route(
    "/users/<int:user_id>/status",
    methods=["GET", "POST"],
)
@admin_required
def change_user_status(user_id):
    user = get_user_by_id(user_id)

    if user is None:
        abort(404)

    # Prevent the logged-in Admin from disabling their own account.
    if user_id == session.get("user_id"):
        return "You cannot change the status of your own account.", 400

    error = None

    if request.method == "POST":
        new_status = (
            "Inactive"
            if user["status"] == "Active"
            else "Active"
        )

        try:
            update_user_status(
                user_id,
                new_status,
            )

            return redirect(url_for("users"))

        except ValueError as error_message:
            error = str(error_message)

        except mysql.connector.Error:
            error = "Unable to change the user status."

    return render_template(
        "change_user_status.html",
        user=user,
        error=error,
    )


@app.route("/dashboard")
@login_required
def dashboard():
    stats = get_dashboard_stats()

    return render_template(
        "dashboard.html",
        stats=stats,
    )

@app.route("/reports")
@login_required
def reports():
    status_report = get_asset_status_report()
    type_report = get_asset_type_report()

    return render_template(
        "reports.html",
        status_report=status_report,
        type_report=type_report,
    )


@app.route("/reports/export")
@login_required
def export_full_report():
    assets = get_assets_for_export()
    employees = get_employees_for_export()
    assignments = get_assignments_for_export()
    maintenance = get_maintenance_for_export()

    workbook = Workbook()

    # Remove the default sheet.
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        bold=True,
        color="FFFFFF",
    )

    def safe_value(value):
        # Convert dates to text so even unusual test dates
        # such as year 0225 will display instead of #######.
        if isinstance(value, datetime):
            return value.strftime("%Y-%m-%d %H:%M:%S")

        if isinstance(value, date):
            return value.isoformat()

        return value

    def setup_sheet(
        title,
        headers,
        rows,
        keys,
    ):
        worksheet = workbook.create_sheet(title)

        worksheet.append(headers)

        for cell in worksheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        for row in rows:
            worksheet.append(
                [
                    safe_value(row[key])
                    for key in keys
                ]
            )

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        # Automatically size the columns.
        for column_cells in worksheet.columns:
            max_length = 0

            column_letter = get_column_letter(
                column_cells[0].column
            )

            for cell in column_cells:
                value = cell.value

                if value is None:
                    continue

                max_length = max(
                    max_length,
                    len(str(value)),
                )

            worksheet.column_dimensions[
                column_letter
            ].width = min(max_length + 2, 40)

        return worksheet

    # -------------------------
    # Assets
    # -------------------------

    setup_sheet(
        "Assets",
        [
            "ID",
            "Asset Tag",
            "Device Name",
            "Asset Type",
            "Brand",
            "Model",
            "Serial Number",
            "Status",
            "Purchase Date",
            "Notes",
            "Created At",
        ],
        assets,
        [
            "asset_id",
            "asset_tag",
            "device_name",
            "asset_type",
            "brand",
            "model",
            "serial_number",
            "status",
            "purchase_date",
            "notes",
            "created_at",
        ],
    )

    # -------------------------
    # Employees
    # -------------------------

    setup_sheet(
        "Employees",
        [
            "ID",
            "Employee Code",
            "Full Name",
            "Department",
            "Position",
            "Email",
            "Status",
            "Created At",
        ],
        employees,
        [
            "employee_id",
            "employee_code",
            "full_name",
            "department",
            "position",
            "email",
            "status",
            "created_at",
        ],
    )

    # -------------------------
    # Assignments
    # -------------------------

    setup_sheet(
        "Assignments",
        [
            "ID",
            "Asset Tag",
            "Device Name",
            "Employee Code",
            "Employee",
            "Department",
            "Assigned Date",
            "Returned Date",
            "Status",
            "Notes",
        ],
        assignments,
        [
            "assignment_id",
            "asset_tag",
            "device_name",
            "employee_code",
            "full_name",
            "department",
            "assigned_date",
            "returned_date",
            "status",
            "notes",
        ],
    )

    # -------------------------
    # Maintenance
    # -------------------------

    maintenance_sheet = setup_sheet(
        "Maintenance",
        [
            "ID",
            "Asset Tag",
            "Device Name",
            "Maintenance Date",
            "Type",
            "Issue",
            "Action Taken",
            "Technician / Vendor",
            "Cost",
            "Status",
            "Completed Date",
            "Notes",
        ],
        maintenance,
        [
            "maintenance_id",
            "asset_tag",
            "device_name",
            "maintenance_date",
            "maintenance_type",
            "issue_description",
            "action_taken",
            "technician_vendor",
            "cost",
            "status",
            "completed_date",
            "notes",
        ],
    )

    # Format Maintenance Cost column.
    for cell in maintenance_sheet["I"][1:]:
        cell.number_format = '₱#,##0.00'

    output = BytesIO()

    workbook.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="IT_Asset_Management_Report.xlsx",
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )

@app.route(
    "/assets/<int:asset_id>/delete",
    methods=["GET", "POST"],
)
@admin_required
def delete_asset_route(asset_id):
    asset = get_asset_by_id(asset_id)

    if asset is None:
        abort(404)

    if asset["status"] != "Retired":
        return "Only retired assets can be permanently deleted.", 400

    has_history = asset_has_history(asset_id)

    if request.method == "POST":
        if has_history:
            return (
                "This asset has assignment or maintenance history "
                "and cannot be permanently deleted.",
                400,
            )

        try:
            delete_asset(asset_id)
            return redirect(url_for("home"))

        except mysql.connector.Error:
            return render_template(
                "delete_asset.html",
                asset=asset,
                has_history=has_history,
                error="Unable to delete the asset. Please try again.",
            )

    return render_template(
        "delete_asset.html",
        asset=asset,
        has_history=has_history,
        error=None,
    )


if __name__ == "__main__":
    app.run(debug=os.getenv("FLASK_DEBUG", "0") == "1", use_reloader=False)
